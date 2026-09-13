from http.server import HTTPServer, SimpleHTTPRequestHandler
import os
import subprocess
import tempfile
import shutil
import mimetypes
import re
from pypdf import PdfReader, PdfWriter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
WEB_DIR = os.path.join(BASE_DIR, "www")

MAX_SIZE = 50 * 1024 * 1024
MAGICK = shutil.which("magick") or "/data/data/com.termux/files/usr/bin/magick"


def send_file(handler, path, content_type, filename):
    with open(path, "rb") as f:
        data = f.read()

    handler.send_response(200)
    handler.send_header("Content-Type", content_type)
    handler.send_header(
        "Content-Disposition",
        f'attachment; filename="{filename}"'
    )
    handler.send_header("Content-Length", str(len(data)))
    handler.end_headers()
    handler.wfile.write(data)


def parse_multipart(handler):
    content_type = handler.headers.get("Content-Type", "")
    match = re.search(r'boundary="?([^";]+)"?', content_type)

    if not match:
        raise ValueError("Invalid multipart request")

    boundary = match.group(1).encode()
    size = int(handler.headers.get("Content-Length", "0"))

    if size <= 0 or size > MAX_SIZE:
        raise ValueError("File too large")

    body = handler.rfile.read(size)
    parts = body.split(b"--" + boundary)

    files = []
    fields = {}

    for part in parts:
        part = part.strip(b"\r\n-")
        if not part or b"\r\n\r\n" not in part:
            continue

        headers, data = part.split(b"\r\n\r\n", 1)
        headers_text = headers.decode("utf-8", "replace")

        disposition = re.search(
            r'Content-Disposition:.*?name="([^"]+)"(?:;\s*filename="([^"]*)")?',
            headers_text,
            re.I
        )

        if not disposition:
            continue

        name = disposition.group(1)
        filename = disposition.group(2)

        if filename is not None:
            files.append((name, os.path.basename(filename), data))
        else:
            fields[name] = data.decode("utf-8", "replace")

    return files, fields


class ToolBoxHandler(SimpleHTTPRequestHandler):

    def do_POST(self):

        if self.path == "/compress":
            self.compress_pdf()
            return

        if self.path == "/jpg-to-pdf":
            self.jpg_to_pdf()
            return

        if self.path == "/pdf-to-jpg":
            self.pdf_to_jpg()
            return

        if self.path == "/split-pdf":
            self.split_pdf()
            return
        if self.path == "/merge-pdf":
            self.merge_pdf()
            return


        if self.path == "/word-to-pdf":
            self.word_to_pdf()
            return

        if self.path == "/excel-to-pdf":
            self.excel_to_pdf()
            return

        if self.path == "/image-compress":
            self.image_compress()
            return

        if self.path == "/image-resize":
            self.image_resize()
            return

        self.send_error(404)
        return

    def compress_pdf(self):
        try:
            size = int(self.headers.get("Content-Length", "0"))

            if size <= 0 or size > MAX_SIZE:
                self.send_error(
                    413,
                    "PDF must be between 1 byte and 50 MB"
                )
                return

            pdf_data = self.rfile.read(size)

            with tempfile.TemporaryDirectory() as tmp:
                input_pdf = os.path.join(tmp, "input.pdf")
                output_pdf = os.path.join(tmp, "compressed.pdf")

                with open(input_pdf, "wb") as f:
                    f.write(pdf_data)

                subprocess.run(
                    [
                        "gs",
                        "-sDEVICE=pdfwrite",
                        "-dCompatibilityLevel=1.4",
                        "-dPDFSETTINGS=/ebook",
                        "-dNOPAUSE",
                        "-dQUIET",
                        "-dBATCH",
                        f"-sOutputFile={output_pdf}",
                        input_pdf,
                    ],
                    check=True,
                )

                send_file(
                    self,
                    output_pdf,
                    "application/pdf",
                    "ToolBox-compressed.pdf"
                )

        except subprocess.CalledProcessError:
            self.send_error(500, "PDF compression failed")
        except Exception as e:
            self.send_error(500, str(e))

    def jpg_to_pdf(self):
        try:
            files, _ = parse_multipart(self)

            images = [
                item for item in files
                if item[1].lower().endswith((".jpg", ".jpeg"))
            ]

            if not images:
                self.send_error(400, "JPG files required")
                return

            with tempfile.TemporaryDirectory() as tmp:
                input_paths = []

                for index, (_, filename, data) in enumerate(images):
                    path = os.path.join(tmp, f"image_{index}.jpg")
                    with open(path, "wb") as f:
                        f.write(data)
                    input_paths.append(path)

                output_pdf = os.path.join(tmp, "ToolBox-images.pdf")

                subprocess.run(
                    [
                        MAGICK,
                        *input_paths,
                        "-auto-orient",
                        output_pdf
                    ],
                    check=True
                )

                send_file(
                    self,
                    output_pdf,
                    "application/pdf",
                    "ToolBox-images.pdf"
                )

        except subprocess.CalledProcessError:
            self.send_error(500, "JPG to PDF conversion failed")
        except Exception as e:
            self.send_error(500, str(e))

    def split_pdf(self):
        try:
            files, fields = parse_multipart(self)

            pdfs = [
                item for item in files
                if item[1].lower().endswith(".pdf")
            ]

            if not pdfs:
                self.send_error(400, "PDF file required")
                return

            page_range = fields.get("pages", "").strip()

            if not page_range:
                self.send_error(400, "Pages required")
                return

            _, filename, data = pdfs[0]

            with tempfile.TemporaryDirectory() as tmp:
                input_pdf = os.path.join(tmp, "input.pdf")
                output_pdf = os.path.join(tmp, "ToolBox-split.pdf")

                with open(input_pdf, "wb") as f:
                    f.write(data)

                reader = PdfReader(input_pdf)
                writer = PdfWriter()
                total_pages = len(reader.pages)

                selected = set()

                for part in page_range.split(","):
                    part = part.strip()

                    if "-" in part:
                        start, end = part.split("-", 1)
                        start = int(start)
                        end = int(end)

                        if start > end:
                            start, end = end, start

                        for page in range(start, end + 1):
                            selected.add(page)
                    else:
                        selected.add(int(part))

                for page_number in sorted(selected):
                    if page_number < 1 or page_number > total_pages:
                        self.send_error(
                            400,
                            f"Page {page_number} is outside the PDF"
                        )
                        return

                    writer.add_page(reader.pages[page_number - 1])

                with open(output_pdf, "wb") as f:
                    writer.write(f)

                send_file(
                    self,
                    output_pdf,
                    "application/pdf",
                    "ToolBox-split.pdf"
                )

        except (ValueError, IndexError):
            self.send_error(400, "Invalid page range")
        except Exception as e:
            self.send_error(500, str(e))

    def merge_pdf(self):
        try:
            files, fields = parse_multipart(self)

            pdfs = [
                item for item in files
                if item[1].lower().endswith(".pdf")
            ]

            if len(pdfs) < 2:
                self.send_error(400, "At least two PDF files required")
                return

            with tempfile.TemporaryDirectory() as tmp:
                output_pdf = os.path.join(tmp, "ToolBox-merged.pdf")
                writer = PdfWriter()

                for _, filename, data in pdfs:
                    input_pdf = os.path.join(
                        tmp,
                        f"{len(writer.pages)}_{filename}"
                    )

                    with open(input_pdf, "wb") as f:
                        f.write(data)

                    reader = PdfReader(input_pdf)

                    for page in reader.pages:
                        writer.add_page(page)

                with open(output_pdf, "wb") as f:
                    writer.write(f)

                send_file(
                    self,
                    output_pdf,
                    "application/pdf",
                    "ToolBox-merged.pdf"
                )

        except Exception as e:
            self.send_error(500, str(e))

    def excel_to_pdf(self):
        try:
            files, _ = parse_multipart(self)

            sheets = [
                item for item in files
                if item[1].lower().endswith((".xlsx", ".xls"))
            ]

            if not sheets:
                self.send_error(400, "Excel file required")
                return

            _, filename, data = sheets[0]

            with tempfile.TemporaryDirectory() as tmp:
                input_path = os.path.join(tmp, "input.xlsx")
                output_pdf = os.path.join(tmp, "ToolBox-excel.pdf")

                with open(input_path, "wb") as f:
                    f.write(data)

                subprocess.run(
                    [
                        "python",
                        "-c",
                        """
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
import os, sys

src = sys.argv[1]
out = sys.argv[2]

wb = load_workbook(src, data_only=True)
styles = getSampleStyleSheet()
story = []

for ws in wb.worksheets:
    story.append(Paragraph(ws.title, styles["Heading2"]))
    story.append(Spacer(1, 8))

    rows = []
    for row in ws.iter_rows(values_only=True):
        values = ["" if v is None else str(v) for v in row]
        if any(values):
            rows.append(values)

    if not rows:
        continue

    max_cols = max(len(r) for r in rows)
    rows = [r + [""] * (max_cols - len(r)) for r in rows]

    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
        ("BACKGROUND", (0,0), (-1,0), colors.lightgrey),
        ("FONTNAME", (0,0), (-1,-1), "Helvetica"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 5),
        ("RIGHTPADDING", (0,0), (-1,-1), 5),
    ]))
    story.append(table)
    story.append(Spacer(1, 18))

doc = SimpleDocTemplate(
    out,
    pagesize=landscape(A4),
    rightMargin=24,
    leftMargin=24,
    topMargin=24,
    bottomMargin=24,
)

doc.build(story)
""",
                        input_path,
                        output_pdf,
                    ],
                    check=True,
                )

                send_file(
                    self,
                    output_pdf,
                    "application/pdf",
                    "ToolBox-excel.pdf"
                )

        except subprocess.CalledProcessError:
            self.send_error(500, "Excel to PDF conversion failed")
        except Exception as e:
            self.send_error(500, str(e))

    def word_to_pdf(self):
        try:
            files, _ = parse_multipart(self)

            docs = [
                item for item in files
                if item[1].lower().endswith(".docx")
            ]

            if not docs:
                self.send_error(400, "DOCX file required")
                return

            _, filename, data = docs[0]

            with tempfile.TemporaryDirectory() as tmp:
                input_path = os.path.join(tmp, "input.docx")
                output_tex = os.path.join(tmp, "ToolBox-word.tex")
                output_pdf = os.path.join(tmp, "ToolBox-word.pdf")

                with open(input_path, "wb") as f:
                    f.write(data)

                subprocess.run(
                    [
                        "pandoc",
                        input_path,
                        "--standalone",
                        "-t",
                        "latex",
                        "-o",
                        output_tex,
                    ],
                    check=True,
                )

                subprocess.run(
                    [
                        "tectonic",
                        output_tex,
                    ],
                    check=True,
                )

                send_file(
                    self,
                    output_pdf,
                    "application/pdf",
                    "ToolBox-word.pdf"
                )

        except subprocess.CalledProcessError:
            self.send_error(500, "Word to PDF conversion failed")
        except Exception as e:
            self.send_error(500, str(e))

    def pdf_to_jpg(self):
        try:
            files, fields = parse_multipart(self)

            pdfs = [
                item for item in files
                if item[1].lower().endswith(".pdf")
            ]

            if not pdfs:
                self.send_error(400, "PDF file required")
                return

            _, filename, data = pdfs[0]

            page_range = fields.get("pages", "").strip()

            with tempfile.TemporaryDirectory() as tmp:
                input_pdf = os.path.join(tmp, "input.pdf")

                with open(input_pdf, "wb") as f:
                    f.write(data)

                reader = PdfReader(input_pdf)
                total_pages = len(reader.pages)

                if not total_pages:
                    self.send_error(400, "PDF contains no pages")
                    return

                selected = []

                if page_range:
                    selected_set = set()

                    for part in page_range.split(","):
                        part = part.strip()

                        if not part:
                            continue

                        if "-" in part:
                            start_page, end_page = part.split("-", 1)
                            start_page = int(start_page)
                            end_page = int(end_page)

                            if start_page > end_page:
                                start_page, end_page = end_page, start_page

                            for page in range(start_page, end_page + 1):
                                selected_set.add(page)
                        else:
                            selected_set.add(int(part))

                    for page in sorted(selected_set):
                        if page < 1 or page > total_pages:
                            self.send_error(
                                400,
                                f"Page {page} is outside the PDF"
                            )
                            return

                    selected = sorted(selected_set)
                else:
                    selected = list(range(1, total_pages + 1))

                selected_pdf = os.path.join(tmp, "selected.pdf")
                writer = PdfWriter()

                for page_number in selected:
                    writer.add_page(reader.pages[page_number - 1])

                with open(selected_pdf, "wb") as f:
                    writer.write(f)

                output_pattern = os.path.join(tmp, "page-%03d.jpg")

                subprocess.run(
                    [
                        MAGICK,
                        "-density", "150",
                        selected_pdf,
                        "-background", "white",
                        "-alpha", "remove",
                        "-alpha", "off",
                        "-quality", "90",
                        output_pattern
                    ],
                    check=True
                )

                jpgs = sorted(
                    os.path.join(tmp, name)
                    for name in os.listdir(tmp)
                    if name.lower().endswith(".jpg")
                )

                if not jpgs:
                    self.send_error(500, "No JPG pages were created")
                    return

                if len(jpgs) == 1:
                    send_file(
                        self,
                        jpgs[0],
                        "image/jpeg",
                        "ToolBox-page-1.jpg"
                    )
                    return

                import zipfile
                archive_path = os.path.join(tmp, "ToolBox-pdf-pages.zip")
                with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_STORED) as z:
                    for jpg in jpgs:
                        z.write(jpg, os.path.basename(jpg))

                send_file(
                    self,
                    archive_path,
                    "application/zip",
                    "ToolBox-pdf-pages.zip"
                )

        except (ValueError, IndexError):
            self.send_error(400, "Invalid page range")
        except subprocess.CalledProcessError:
            self.send_error(500, "PDF to JPG conversion failed")
        except Exception as e:
            self.send_error(500, str(e))

    def image_compress(self):
        try:
            files, fields = parse_multipart(self)

            if not files:
                self.send_error(400, "Image required")
                return

            _, filename, data = files[0]

            ext = os.path.splitext(filename)[1].lower()

            if ext not in (".jpg", ".jpeg", ".png", ".webp"):
                self.send_error(400, "Unsupported image format")
                return

            quality = fields.get("quality", "82")

            try:
                quality = max(20, min(95, int(quality)))
            except ValueError:
                quality = 82

            with tempfile.TemporaryDirectory() as tmp:
                input_path = os.path.join(tmp, "input" + ext)
                output_ext = ".jpg" if ext in (".jpg", ".jpeg") else ext
                output_path = os.path.join(tmp, "compressed" + output_ext)

                with open(input_path, "wb") as f:
                    f.write(data)

                subprocess.run(
                    [
                        MAGICK,
                        input_path,
                        "-auto-orient",
                        "-strip",
                        "-quality",
                        str(quality),
                        output_path
                    ],
                    check=True
                )

                content_type = mimetypes.guess_type(output_path)[0] or "application/octet-stream"

                send_file(
                    self,
                    output_path,
                    content_type,
                    "ToolBox-compressed" + output_ext
                )

        except subprocess.CalledProcessError:
            self.send_error(500, "Image compression failed")
        except Exception as e:
            self.send_error(500, str(e))

    def image_resize(self):
        try:
            files, fields = parse_multipart(self)

            if not files:
                self.send_error(400, "Image required")
                return

            _, filename, data = files[0]

            ext = os.path.splitext(filename)[1].lower()

            if ext not in (".jpg", ".jpeg", ".png", ".webp"):
                self.send_error(400, "Unsupported image format")
                return

            width = int(fields.get("width", "0"))
            height = int(fields.get("height", "0"))
            keep_ratio = fields.get("keep_ratio", "true") == "true"

            if width < 1 or width > 10000 or height < 1 or height > 10000:
                self.send_error(400, "Invalid dimensions")
                return

            with tempfile.TemporaryDirectory() as tmp:
                input_path = os.path.join(tmp, "input" + ext)
                output_path = os.path.join(tmp, "resized" + ext)

                with open(input_path, "wb") as f:
                    f.write(data)

                geometry = f"{width}x{height}"

                if not keep_ratio:
                    geometry += "!"

                subprocess.run(
                    [
                        MAGICK,
                        input_path,
                        "-auto-orient",
                        "-resize",
                        geometry,
                        output_path
                    ],
                    check=True
                )

                content_type = mimetypes.guess_type(output_path)[0] or "application/octet-stream"

                send_file(
                    self,
                    output_path,
                    content_type,
                    "ToolBox-resized" + ext
                )

        except (ValueError, TypeError):
            self.send_error(400, "Invalid resize values")
        except subprocess.CalledProcessError:
            self.send_error(500, "Image resize failed")
        except Exception as e:
            self.send_error(500, str(e))


os.chdir(WEB_DIR)

server = HTTPServer(("0.0.0.0", 8081), ToolBoxHandler)

print("ToolBox running on http://127.0.0.1:8081")

server.serve_forever()
